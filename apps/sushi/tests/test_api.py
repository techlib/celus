from io import BytesIO
from unittest import mock

import pandas as pd
import pytest
from django.urls import reverse
from openpyxl import load_workbook
from scheduler.fake_data import FetchIntentionFactory
from scheduler.models import FetchIntention

from sushi.fake_data import FetchAttemptFactory
from sushi.models import SushiCredentials
from test_scenarios.basic import *  # noqa


@pytest.mark.django_db
class TestSushiCredentialsAPI:
    @pytest.mark.django_db(transaction=True)
    @pytest.mark.parametrize(
        [
            "user",
            "delete_credentials",
            "delete_fetchintentions",
            "delete_data",
            "delete_fetchattempts_and_related_importbatches",
            "status_code",
        ],
        [
            ["admin2", True, True, True, True, 204],
            ["admin1", False, False, True, False, 404],
            ["user2", False, False, True, False, 404],
            ["admin2", True, True, False, False, 204],
            ["admin1", False, False, False, False, 404],
        ],
    )
    def test_destroy(
        self,
        credentials,
        clients,
        user,
        basic1,
        delete_fetchattempts_and_related_importbatches,
        delete_credentials,
        delete_fetchintentions,
        delete_data,
        status_code,
    ):
        cr = credentials["standalone_tr"]
        fetch_attempts = FetchAttemptFactory.create_batch(2, credentials=cr)
        fi = FetchIntentionFactory(credentials=cr, attempt=fetch_attempts[0])

        cr_queryset = SushiCredentials.objects.filter(pk=cr.pk)
        fi_queryset = FetchIntention.objects.filter(pk=fi.pk)

        assert cr_queryset.exists()
        assert fi_queryset.exists()

        with mock.patch(
            "sushi.views.delete_fetchattempts_and_related_importbatches_task"
        ) as mock_task:
            url = reverse("sushi-credentials-detail", args=[cr.pk])
            url += f"?delete_data={delete_data}"
            res = clients[user].delete(url)
            assert res.status_code == status_code

            if delete_fetchattempts_and_related_importbatches:
                fa_pks = [fa.pk for fa in fetch_attempts]
                mock_task.delay.assert_called_with(fa_pks)
            else:
                mock_task.delay.assert_not_called()

        if delete_credentials:
            assert not cr_queryset.exists()
        else:
            assert cr_queryset.exists()

        if delete_fetchintentions:
            assert not fi_queryset.exists()
        else:
            assert fi_queryset.exists()


def get_core_attrs_credentials():
    return {
        "title": None,
        "organization": None,
        "publisher/vendor/platform": None,
        "SUSHI url": None,
        "requestor id": None,
        "customer id": None,
    }


def get_empty_credentials_c4():
    empty_credentials_c4 = get_core_attrs_credentials()
    empty_credentials_c4.update(
        {
            key: None
            for key in [
                "http username",
                "http password",
                "extra params",
                "BR1",
                "BR2",
                "BR3",
                "DB1",
                "DB2",
                "JR1",
                "JR1a",
                "JR1GOA",
                "JR2",
                "JR5",
                "MR1",
                "PR1",
            ]
        }
    )
    return empty_credentials_c4


def get_empty_credentials_c5():
    empty_credentials_c5 = get_core_attrs_credentials()
    empty_credentials_c5.update(
        {key: None for key in ["api key", "platform filter", "DR", "IR", "PR", "TR"]}
    )
    return empty_credentials_c5


def get_empty_credentials_c51():
    empty_credentials_c51 = get_core_attrs_credentials()
    empty_credentials_c51.update(
        {key: None for key in ["api key", "platform filter", "DR", "IR", "PR", "TR"]}
    )
    return empty_credentials_c51


@pytest.fixture
def credentials_dict(credentials):
    standalone_br1_jr1 = get_empty_credentials_c4()
    cr = credentials["standalone_br1_jr1"]
    standalone_br1_jr1.update(
        {
            "title": cr.title,
            "organization": cr.organization.name_en,
            "publisher/vendor/platform": cr.platform.name_en,
            "SUSHI url": cr.url,
            "requestor id": cr.requestor_id,
            "customer id": cr.customer_id,
            "http username": cr.http_username,
            "http password": cr.http_password,
            "extra params": str(cr.extra_params),
            "BR1": "active",
            "JR1": "active",
        }
    )

    branch_pr = get_empty_credentials_c5()
    cr = credentials["branch_pr"]
    branch_pr.update(
        {
            "title": cr.title,
            "organization": cr.organization.name_en,
            "publisher/vendor/platform": cr.platform.name_en,
            "SUSHI url": cr.url,
            "requestor id": cr.requestor_id,
            "customer id": cr.customer_id,
            "api key": cr.api_key,
            "platform filter": str(cr.extra_params["platform"])
            if "platform" in cr.extra_params
            else None,
            "PR": "active",
        }
    )

    standalone_tr = get_empty_credentials_c5()
    cr = credentials["standalone_tr"]
    standalone_tr.update(
        {
            "title": cr.title,
            "organization": cr.organization.name_en,
            "publisher/vendor/platform": cr.platform.name_en,
            "SUSHI url": cr.url,
            "requestor id": cr.requestor_id,
            "customer id": cr.customer_id,
            "api key": cr.api_key,
            "platform filter": str(cr.extra_params["platform"])
            if "platform" in cr.extra_params
            else None,
            "TR": "active",
        }
    )

    standalone_ir51 = get_empty_credentials_c51()
    cr = credentials["standalone_ir51"]
    standalone_ir51.update(
        {
            "title": cr.title,
            "organization": cr.organization.name_en,
            "publisher/vendor/platform": cr.platform.name_en,
            "SUSHI url": cr.url,
            "requestor id": cr.requestor_id,
            "customer id": cr.customer_id,
            "api key": cr.api_key,
            "platform filter": str(cr.extra_params["platform"])
            if "platform" in cr.extra_params
            else None,
            "IR": "active",
        }
    )

    del credentials
    return locals()


@pytest.fixture
def sushi_cred_dataframe_fixture(credentials_dict):
    su_all = {
        "Credentials-COUNTER4": pd.DataFrame(credentials_dict["standalone_br1_jr1"], index=[0]),
        "Credentials-COUNTER5": pd.DataFrame(
            [credentials_dict["branch_pr"], credentials_dict["standalone_tr"]]
        ),
        "Credentials-COUNTER5.1": pd.DataFrame(credentials_dict["standalone_ir51"], index=[0]),
    }
    su_standalone = {
        "Credentials-COUNTER4": pd.DataFrame(credentials_dict["standalone_br1_jr1"], index=[0]),
        "Credentials-COUNTER5": pd.DataFrame(credentials_dict["standalone_tr"], index=[0]),
        "Credentials-COUNTER5.1": pd.DataFrame(credentials_dict["standalone_ir51"], index=[0]),
    }
    admin1_root = {
        "Credentials-COUNTER4": pd.DataFrame(get_empty_credentials_c4(), index=[]),
        "Credentials-COUNTER5": pd.DataFrame(get_empty_credentials_c5(), index=[]),
        "Credentials-COUNTER5.1": pd.DataFrame(get_empty_credentials_c51(), index=[]),
    }
    admin2_standalone = {
        "Credentials-COUNTER4": pd.DataFrame(credentials_dict["standalone_br1_jr1"], index=[0]),
        "Credentials-COUNTER5": pd.DataFrame(credentials_dict["standalone_tr"], index=[0]),
        "Credentials-COUNTER5.1": pd.DataFrame(credentials_dict["standalone_ir51"], index=[0]),
    }
    return locals()


@pytest.fixture
def sushi_cred_with_platforms_dataframe_fixture(sushi_cred_dataframe_fixture):
    def sort(df, all_organizations) -> pd.DataFrame:
        to_sort_by = ["publisher/vendor/platform", "title"]
        if all_organizations:
            to_sort_by.insert(1, "organization")
        return df.sort_values(by=to_sort_by, ignore_index=True)

    def add_platforms_to_dataframe(df, platforms):
        new_platform_df = pd.DataFrame({"publisher/vendor/platform": platforms})
        df = pd.concat([df, new_platform_df])
        return df

    def get_the_dict_of_dataframes(df_map, platforms, all_organizations=False, organizations=None):
        dict_of_dataframes = {}
        for sheetname in ["Credentials-COUNTER5", "Credentials-COUNTER5.1"]:
            used_platforms = {e[2] for e in df_map[sheetname].to_records()}
            df = add_platforms_to_dataframe(
                df_map[sheetname], [e for e in platforms if e not in used_platforms]
            )
            df = df.drop(columns=["SUSHI url", "TR", "DR", "PR", "IR"], errors="ignore")
            if not all_organizations:
                df.drop(columns=["organization"], inplace=True)
            df = sort(df, all_organizations)
            dict_of_dataframes[sheetname] = df
        if all_organizations:
            dict_of_dataframes["Organizations"] = pd.DataFrame({"organization": organizations})
        return dict_of_dataframes

    su_all = get_the_dict_of_dataframes(
        sushi_cred_dataframe_fixture["su_all"],
        platforms=["root", "shared", "brain", "branch"],
        all_organizations=True,
        organizations=["empty", "master", "root", "branch", "standalone"],
    )

    su_standalone = get_the_dict_of_dataframes(
        sushi_cred_dataframe_fixture["su_standalone"],
        platforms=["shared", "brain"],
        all_organizations=False,
    )

    admin1_root = get_the_dict_of_dataframes(
        sushi_cred_dataframe_fixture["admin1_root"],
        platforms=["root", "shared", "brain"],
        all_organizations=False,
    )

    admin2_standalone = get_the_dict_of_dataframes(
        sushi_cred_dataframe_fixture["admin2_standalone"],
        platforms=["shared", "brain"],
        all_organizations=False,
    )

    return {
        "su_all": su_all,
        "su_standalone": su_standalone,
        "admin1_root": admin1_root,
        "admin2_standalone": admin2_standalone,
    }


@pytest.mark.django_db
class TestSushiCredentialsExport:
    @pytest.mark.parametrize(
        ["identity", "organization", "status_code"],
        [
            ["su", "standalone", 200],
            ["su", "branch", 200],
            ["su", None, 200],
            ["admin1", None, 404],
            ["admin1", "standalone", 404],
            ["admin1", "branch", 200],
            ["admin2", None, 404],
            ["admin2", "standalone", 200],
            ["admin2", "branch", 404],
            ["user1", None, 404],
            ["user1", "standalone", 404],
            ["user1", "branch", 404],
            ["user2", None, 404],
            ["user2", "standalone", 404],
            ["user2", "branch", 404],
        ],
    )
    def test_user_permission_for_credentials_export(
        self, clients, identity, basic1, credentials, organization, organizations, status_code
    ):
        organization_id = organizations[organization].pk if organization else -1
        url = reverse("sushi-credentials-export-credentials")
        url += f"?organization={organization_id}"
        resp = clients[identity].post(url, data={"pk": [x.pk for x in credentials.values()]})
        assert resp.status_code == status_code

    @pytest.mark.parametrize(
        ["identity", "organization", "case"],
        [
            ["su", None, "su_all"],
            ["su", "standalone", "su_standalone"],
            ["admin1", "root", "admin1_root"],
            ["admin2", "standalone", "admin2_standalone"],
        ],
    )
    def test_dataframes_in_credentials_export(
        self,
        identity,
        clients,
        basic1,
        credentials,
        sushi_cred_dataframe_fixture,
        organizations,
        organization,
        case,
    ):
        organization_id = organizations[organization].pk if organization else -1
        url = reverse("sushi-credentials-export-credentials")
        url += f"?organization={organization_id}"
        resp = clients[identity].post(url, data={"pk": [x.pk for x in credentials.values()]})
        resp_workbook = load_workbook(filename=BytesIO(resp.content), read_only=True)
        sheetnames = {"Credentials-COUNTER4", "Credentials-COUNTER5", "Credentials-COUNTER5.1"}
        assert sheetnames == set(resp_workbook.sheetnames)
        for sheetname in sheetnames:
            if sheetname not in sushi_cred_dataframe_fixture[case]:
                # Missing the kind of credentials in the list
                # => skip
                continue
            ws = resp_workbook[sheetname]
            data = list(ws.values)[1:]
            resp_df = pd.DataFrame(data=data, columns=[cell.value for cell in ws[1]])
            assert sushi_cred_dataframe_fixture[case][sheetname].equals(resp_df)

    @pytest.mark.parametrize(
        ["identity", "organization", "status_code"],
        [
            ["su", "standalone", 200],
            ["su", "branch", 200],
            ["su", None, 200],
            ["admin1", None, 404],
            ["admin1", "standalone", 404],
            ["admin1", "branch", 200],
            ["admin2", None, 404],
            ["admin2", "standalone", 200],
            ["admin2", "branch", 404],
            ["user1", None, 404],
            ["user1", "standalone", 404],
            ["user1", "branch", 404],
            ["user2", None, 404],
            ["user2", "standalone", 404],
            ["user2", "branch", 404],
        ],
    )
    def test_user_permission_for_credentials_import_template(
        self, clients, identity, basic1, organization, organizations, status_code
    ):
        organization_id = organizations[organization].pk if organization else -1
        url = reverse("sushi-credentials-import-template")
        url += f"?organization={organization_id}"
        resp = clients[identity].get(url)
        assert resp.status_code == status_code

    @pytest.mark.parametrize(
        ["identity", "organization", "case", "sheetnames"],
        [
            [
                "su",
                None,
                "su_all",
                ["Explanation", "Credentials-COUNTER5", "Credentials-COUNTER5.1", "Organizations"],
            ],
            [
                "su",
                "standalone",
                "su_standalone",
                ["Explanation", "Credentials-COUNTER5", "Credentials-COUNTER5.1"],
            ],
            [
                "admin1",
                "root",
                "admin1_root",
                ["Explanation", "Credentials-COUNTER5", "Credentials-COUNTER5.1"],
            ],
            [
                "admin2",
                "standalone",
                "admin2_standalone",
                ["Explanation", "Credentials-COUNTER5", "Credentials-COUNTER5.1"],
            ],
        ],
    )
    def test_dataframes_in_credentials_import_template(
        self,
        identity,
        clients,
        basic1,
        organizations,
        organization,
        case,
        sushi_cred_with_platforms_dataframe_fixture,
        sheetnames,
    ):
        organization_id = organizations[organization].pk if organization else -1
        url = reverse("sushi-credentials-import-template")
        resp = clients[identity].get(url, {"organization": organization_id})
        resp_workbook = load_workbook(filename=BytesIO(resp.content), read_only=True)
        assert set(sheetnames) == set(resp_workbook.sheetnames)
        for sheetname in sheetnames:
            if sheetname == "Explanation":
                continue
            ws = resp_workbook[sheetname]
            data = list(ws.values)[1:]
            resp_df = pd.DataFrame(data=data, columns=[cell.value for cell in ws[1]])
            cols_to_drop = [col for col in resp_df.columns if col is None]
            resp_df.drop(cols_to_drop, axis=1, inplace=True)
            resp_df.dropna(axis=0, how="all", inplace=True)
            assert sushi_cred_with_platforms_dataframe_fixture[case][sheetname].equals(resp_df)
