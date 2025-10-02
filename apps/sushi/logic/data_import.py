import json
import logging
from collections import Counter
from datetime import timedelta
from enum import Enum, auto
from typing import Optional

import reversion
from core.logic.dates import last_month, month_end, month_start
from django.db.models import Count, Q
from django.utils.timezone import now
from organizations.models import Organization
from publications.logic.knowledgebase import (
    get_provider_for_counter_version,
    is_report_type_whitelisted,
)
from publications.models import Platform
from scheduler.models import FetchIntention, Harvest

from sushi.logic.export import Col

from ..models import CounterReportsToCredentials, CounterReportType, SushiCredentials

logger = logging.getLogger(__name__)


class Perform(Enum):
    UPDATE_NONE = auto()
    UPDATE_NOT_VERIFIED = auto()
    UPDATE_ALL = auto()


def import_sushi_credentials_from_xlsx(
    filename,
    sheet_no: int = 2,
    single_org: Optional[str] = None,
    log_diff: bool = True,
    log_trivial: bool = False,
    update_credentials=Perform.UPDATE_NONE,
    reversion_comment: Optional[str] = None,
    harvest_months: Optional[int] = None,
) -> dict:
    from openpyxl import load_workbook  # noqa - slow import

    workbook = load_workbook(filename=filename, read_only=True)
    if sheet_no > len(workbook.worksheets):
        raise ValueError("chosen sheet doesn't exist")
    credentials_sheet = workbook.worksheets[sheet_no - 1]
    # Determine whether the credentials are C5 or C51 based on the sheet name
    counter_version = 51 if credentials_sheet.title == "Credentials-COUNTER5.1" else 5
    headers = [header.value for header in credentials_sheet[1] if header.value]
    if Col.PUBLISHER_VENDOR_PLATFORM.value not in headers or Col.CUSTOMER_ID.value not in headers:
        raise ValueError("essential headers are missing")

    EXPECTED_COLS = {
        Col.TITLE,
        Col.PUBLISHER_VENDOR_PLATFORM,
        Col.REQUESTOR_ID,
        Col.CUSTOMER_ID,
        Col.API_KEY,
        Col.PLATFORM_FILTER,
    }
    if single_org:
        if isinstance(single_org, int) or single_org.isdigit():
            single_org = Organization.objects.get(pk=single_org)
        else:
            single_org = Organization.objects.get(Q(name_en=single_org) | Q(short_name=single_org))

    elif Col.ORGANIZATION.value not in headers:
        raise ValueError(
            "provide correct organization reference if importing for a single organization "
            "or add 'organization' column to the sheet if importing for consortium"
        )
    else:
        EXPECTED_COLS.add(Col.ORGANIZATION)

    # checking expected headers
    EXPECTED_COLS = {col.value for col in EXPECTED_COLS}
    missing_cols = EXPECTED_COLS - set(headers)
    if missing_cols:
        logger.warning(
            f"Following columns are missing in the sheet: {', '.join(missing_cols)}. "
            "Make sure you don't need to import data expected to be in these columns."
        )
    excessive_cols = set(headers) - EXPECTED_COLS
    if excessive_cols:
        logger.warning(
            f"Following columns are not expected in the sheet: {', '.join(excessive_cols)}. "
            "Data in these columns will be ignored."
        )
    start_row = 2
    end_row = credentials_sheet.max_row
    rows = credentials_sheet.iter_rows(min_row=start_row, max_row=end_row, values_only=True)
    records = [dict(zip(headers, row)) for row in rows]

    return import_sushi_credentials_new(
        records,
        single_org=single_org,
        log_diff=log_diff,
        log_trivial=log_trivial,
        update_credentials=update_credentials,
        reversion_comment=reversion_comment,
        counter_version=counter_version,
        harvest_months=harvest_months,
    )


def import_sushi_credentials_new(
    records: [dict],
    single_org: Optional[Organization] = None,
    log_diff: bool = True,
    log_trivial: bool = False,
    update_credentials=Perform.UPDATE_NONE,
    reversion_comment: Optional[str] = None,
    counter_version: int = 5,
    harvest_months: Optional[int] = None,
) -> dict:
    """
    Imports SUSHI credentials from a list of dicts describing the data - new version for xlsx
    :param update_credentials: method of deciding which credentials should be updated
    :param log_trivial: log also trivial messages (e.g. skipped empty rows)
    :param log_diff: log differences between db and imported data
    :param single_org: if provided, this organization will be used for all imported credentials
           otherwise, organization will be taken from the data
    :param reversion_comment: comment that will be passed to the reversion version, if None a
           default will be provided
    :param counter_version: which counter version is used for these credentials
    :param records:
    :return:
    """
    stats = Counter()
    sushi_credentials = SushiCredentials.objects.all()
    db_identical_credentials = set(
        sushi_credentials.values_list("organization_id", "platform_id", "counter_version")
        .alias(count=Count("id"))
        .filter(count__gt=1)
    )

    db_credentials = {
        (cr.organization_id, cr.platform_id, cr.counter_version): cr for cr in sushi_credentials
    }
    platform_objects = Platform.objects.all()
    source_id = lambda pl: pl.source.organization_id if pl.source else None  # noqa: E731
    db_platforms = {((pl.name_en or pl.name).lower(), source_id(pl)): pl for pl in platform_objects}

    organization_objects = Organization.objects.all()
    db_organizations = {org.internal_id: org for org in organization_objects}
    db_organizations.update(
        {org.name_en.lower(): org for org in organization_objects if org.name_en is not None}
    )

    def to_clean_str(value):
        if value is None:
            return value
        if not isinstance(value, str):
            value = str(value)
        return " ".join(value.split())

    def log(message: str, *args, stat_name="error", trivial=False, level=logging.ERROR):
        if log_trivial or not trivial:
            logger.log(level, f"row #%03d: {message}", i + 2, *args)
        stats[stat_name] += 1

    new_credentials = []
    for i, record in enumerate(records):  # noqa: B007 - i is used in `log` function
        customer_id = to_clean_str(record.get(Col.CUSTOMER_ID.value))
        if not customer_id:
            log(
                "Customer ID empty - interpreting as empty row",
                stat_name="empty_customer_id",
                trivial=True,
                level=logging.INFO,
            )
            # skip empty lines
            continue

        # organization name
        organization_name = (
            single_org.name_en if single_org else to_clean_str(record.get(Col.ORGANIZATION.value))
        )
        if not organization_name:
            log("Organization name empty")
            continue
        organization = db_organizations.get(organization_name.lower())
        if not organization:
            log('Unknown organization: "%s"', organization_name)
            continue

        # platform
        platform_name = to_clean_str(record.get(Col.PUBLISHER_VENDOR_PLATFORM.value))
        plat_source_org: Platform = db_platforms.get((platform_name.lower(), organization.id))
        plat_source_other = db_platforms.get((platform_name.lower(), None))
        if plat_source_org and plat_source_other:
            log(
                'name_en of platform "%s" created by organization "%s" '
                "conflicts with name_en of one of global platforms",
                plat_source_org.name_en,
                organization.name_en,
            )
            continue
        if not plat_source_other and not plat_source_org:
            log('Unknown platform "%s" for organization "%s"', platform_name, organization_name)
            continue
        platform = plat_source_other or plat_source_org

        # optional fields
        optional = {}
        if title := to_clean_str(record.get(Col.TITLE.value)):
            optional["title"] = title
        else:
            optional["title"] = platform.short_name
            if counter_version == 51:
                optional["title"] += " (C51)"

        if platform_filter := to_clean_str(record.get(Col.PLATFORM_FILTER.value)):
            optional["extra_params"] = {"platform": platform_filter}

        if api_key := to_clean_str(record.get(Col.API_KEY.value)):
            optional["api_key"] = api_key

        if requestor_id := to_clean_str(record.get(Col.REQUESTOR_ID.value)):
            optional["requestor_id"] = requestor_id
        provider = None
        if platform.knowledgebase:
            provider = get_provider_for_counter_version(platform.knowledgebase, counter_version)
        if provider and provider.get("provider", {}).get("url"):
            url = provider["provider"]["url"]
        else:
            log("can't assign url due to missing provider for the platform: '%s'", platform.name_en)
            continue

        # Only https urls are stored within Celus so we need to convert it here
        url = url.replace("http://", "https://", 1)

        # sync credentials
        key = (organization.pk, platform.pk, counter_version)
        if key in db_credentials:
            if key in db_identical_credentials:
                log(
                    'Credentials for organization "%s" platform "%s" counter %s: '
                    "have more than one corresponding instance in the database.",
                    organization.name_en,
                    platform.name_en,
                    counter_version,
                    stat_name="duplicates_skipped",
                )
                continue

            to_sync = dict(customer_id=customer_id, url=url, **optional)
            cr = db_credentials[key]
            diff = {}
            for key, value in to_sync.items():
                current_value = getattr(cr, key)
                if value != current_value:
                    setattr(cr, key, value)
                    diff[key] = (current_value, value)
            if diff:
                if (update_credentials == Perform.UPDATE_NONE) or (
                    update_credentials == Perform.UPDATE_NOT_VERIFIED and cr.is_verified
                ):
                    if log_diff:
                        log(
                            _create_diff_info("diff_skipped", cr, organization, platform, diff),
                            stat_name="diff_skipped",
                            level=logging.WARNING,
                        )
                    continue
                with reversion.create_revision():
                    cr.save()
                    reversion.set_comment(
                        reversion_comment
                        or "Updated from logic.data_import.import_sushi_credentials"
                    )
                if log_diff:
                    log(
                        _create_diff_info("diff_updated", cr, organization, platform, diff),
                        stat_name="diff_updated",
                        level=logging.WARNING,
                    )
            else:
                log(
                    "Credentials already exist in the same version",
                    stat_name="skipped",
                    trivial=True,
                    level=logging.INFO,
                )
        else:
            with reversion.create_revision():
                cr = SushiCredentials.objects.create(
                    organization=organization,
                    platform=platform,
                    counter_version=counter_version,
                    customer_id=customer_id,
                    url=url,
                    **optional,
                )
                reversion.set_comment(
                    reversion_comment or "Created by logic.data_import.import_sushi_credentials"
                )
                db_credentials[key] = cr
                new_credentials.append(cr)
            log("Credentials created", stat_name="added", level=logging.WARNING)

        # report type assignment
        linked_rts = {rt.code for rt in cr.counter_reports.all()}
        if provider and provider.get("assigned_report_types"):
            for report_type in provider["assigned_report_types"]:
                if rt := CounterReportType.objects.filter(
                    code=report_type["report_type"], counter_version=counter_version
                ).first():
                    if rt.requires_whitelisting and not is_report_type_whitelisted(
                        provider, rt.code
                    ):
                        log(
                            f"Report type {rt.code} is not whitelisted for platform",
                            stat_name="report_type_not_whitelisted",
                            level=logging.WARNING,
                        )
                        continue
                    if report_type["report_type"] not in linked_rts:
                        CounterReportsToCredentials.objects.create(
                            credentials=cr, counter_report=rt
                        )
                        log(
                            f"Report type {report_type['report_type']} assigned",
                            stat_name="report_type_assigned",
                            level=logging.INFO,
                        )
                else:
                    log(
                        f"Report type {report_type['report_type']} not found",
                        stat_name="report_type_not_found",
                        level=logging.WARNING,
                    )
        else:
            log(
                "No report types assigned to the platform '%s' - no knowledgebase provider",
                platform.name_en,
                stat_name="report_type_not_assigned",
                level=logging.WARNING,
            )
    if harvest_months:
        for cr in new_credentials:
            intentions = []
            for rt in cr.counter_reports.all():
                month = last_month()
                for _i in range(harvest_months):
                    intentions.append(
                        FetchIntention(
                            credentials=cr,
                            start_date=month,
                            end_date=month_end(month),
                            counter_report=rt,
                            not_before=now(),
                        )
                    )
                    month = month_start(month - timedelta(days=16))
            Harvest.plan_harvesting(
                intentions=intentions, priority=FetchIntention.PRIORITY_NOW, user=None
            )
            log(
                "Planned harvest for %d intentions for credentials %d",
                len(intentions),
                cr.pk,
                stat_name="planned_harvest",
                level=logging.INFO,
            )
    return stats


def _create_diff_info(status, cr, organization, platform, diff):
    info = (
        f"{status}: credentials (id: {cr.pk}, organization: {organization.name_en}, "
        f"platform: {platform.name_en}):"
    )
    for key, (current_value, value) in diff.items():
        info += f"\n          {key}: '{current_value}' -> '{value}'"
    return info


def import_sushi_credentials_from_csv(
    filename,
    prefer_knowledgebase_urls: bool = False,
    reversion_comment: Optional[str] = None,
    override_organization: Optional[Organization] = None,
) -> dict:
    from nibbler.logic.dict_reader import get_dict_reader_from_csv  # noqa - slow import

    reader = get_dict_reader_from_csv(filename)
    records = list(reader)  # read all records from the reader
    return import_sushi_credentials_old(
        records,
        prefer_knowledgebase_urls=prefer_knowledgebase_urls,
        reversion_comment=reversion_comment,
        override_organization=override_organization,
    )


def import_sushi_credentials_old(
    records: [dict],
    prefer_knowledgebase_urls: bool = False,
    reversion_comment: Optional[str] = None,
    default_version=5,
    override_organization: Optional[Organization] = None,
) -> dict:
    """
    Imports SUSHI credentials from a list of dicts describing the data
    :param reversion_comment: comment that will be passed to the reversion version, if None a
         default will be provided
    :param prefer_knowledgebase_urls: if True, the urls from the knowledgebase will be used instead
            of the ones from the file
    :param records:
    :param default_version:
    :param override_organization: if provided, this organization will be used for all imported
               credentials otherwise, organization will be taken from the data
    :return:
    """
    stats = Counter()
    db_credentials = {
        (cr.organization_id, cr.platform_id, cr.counter_version): cr
        for cr in SushiCredentials.objects.all()
    }
    platform_objects = Platform.objects.all()
    source_id = lambda pl: pl.source.organization_id if pl.source else None  # noqa: E731
    platforms = {(pl.short_name.lower(), source_id(pl)): pl for pl in platform_objects}
    platforms.update({(pl.name.lower(), source_id(pl)): pl for pl in platform_objects})
    # get organization
    organizations = {}
    if not override_organization:
        organization_objects = Organization.objects.all()
        organizations = {org.internal_id: org for org in organization_objects}
        organizations.update({org.short_name.lower(): org for org in organization_objects})
        organizations.update({org.name.lower(): org for org in organization_objects})

    seen_keys = set()
    for i, record in enumerate(records):
        if override_organization:
            organization = override_organization
        else:
            organization_name = record.get("organization")
            if not organization_name:
                logger.error("#%03d: Organization name is missing", i + 2)
                stats["error"] += 1
                continue
            organization = organizations.get(organization_name.strip().lower())
            if not organization:
                logger.error('#%03d: Unknown organization: "%s"', i + 2, record.get("organization"))
                stats["error"] += 1
                continue
        # at first try global platforms
        platform = platforms.get((record.get("platform").strip().lower(), None)) or platforms.get(
            (record.get("platform").strip().lower(), organization.id)
        )
        if not platform:
            logger.error(
                '#%03d: Unknown platform: "%s" for organization "%s"',
                i + 2,
                record.get("platform", "").strip(),
                organization.short_name,
            )
            stats["error"] += 1
            continue
        # counter version
        if not (version := get_int_value(record, "version")):
            version = get_int_value(record, "counter_version")
        if not version:
            version = default_version
            logger.warning("#%03d: Version not specified, assuming %d", i + 2, version)
        # other stuff
        key = (organization.pk, platform.pk, version)
        if key in seen_keys:
            logger.error(
                '#%03d: Credentials for organization "%s", platform "%s", counter %d: '
                "have more than one corresponding instance in the file. Skipping.",
                i + 2,
                organization.name_en,
                platform.name_en,
                version,
            )
            stats["error"] += 1
            continue
        seen_keys.add(key)
        # extra attrs are in the format: name=value;name=value;...
        extra_attrs = record.get("extra_attrs", {})
        if extra_attrs:
            extra_attrs = parse_params(extra_attrs, version=version)
        # extra params are in json format
        extra_params = record.get("extra_params", {})
        if extra_params:
            extra_attrs.update(json.loads(extra_params))

        optional = {}
        if "auth" in extra_attrs:
            optional["http_username"], optional["http_password"] = extra_attrs["auth"]
            del extra_attrs["auth"]
        else:
            optional["http_username"] = ""
            optional["http_password"] = ""
        if "api_key" in extra_attrs:
            optional["api_key"] = extra_attrs["api_key"]
            del extra_attrs["api_key"]
        elif "api_key" in record:
            optional["api_key"] = record["api_key"]
        else:
            optional["api_key"] = ""
        if "title" in record:
            optional["title"] = record["title"].strip()
        url = record.get("URL") or record.get("url")
        if prefer_knowledgebase_urls:
            if platform.knowledgebase:
                providers = [
                    p
                    for p in platform.knowledgebase.get("providers", [])
                    if p["counter_version"] == version
                    and "provider" in p
                    and "url" in p["provider"]
                ]
                if providers:
                    url = providers[0]["provider"]["url"]
                    stats["url_knowledgebase"] += 1
                else:
                    stats["url_no_provider"] += 1
            else:
                stats["url_no_knowlegdebase"] += 1

        # Only https urls are stored within Celus so we need to convert it here
        url = url.replace("http://", "https://", 1)

        if key in db_credentials:
            # we update it
            cr = db_credentials[key]
            to_sync = dict(
                customer_id=record.get("customer_id"),
                requestor_id=record.get("requestor_id"),
                url=url,
                extra_params=extra_attrs,
                **optional,
            )
            save = False
            for key, value in to_sync.items():
                if value != getattr(cr, key):
                    logger.info(
                        "#%03d: %s changed from '%s' to '%s'", i + 2, key, getattr(cr, key), value
                    )
                    setattr(cr, key, value)
                    save = True

            if save:
                with reversion.create_revision():
                    cr.save()
                    reversion.set_comment(
                        reversion_comment
                        or "Updated from logic.data_import.import_sushi_credentials"
                    )
                stats["synced"] += 1
            else:
                stats["skipped"] += 1
        else:
            with reversion.create_revision():
                cr = SushiCredentials.objects.create(
                    organization=organization,
                    platform=platform,
                    counter_version=version,
                    customer_id=record.get("customer_id"),
                    requestor_id=record.get("requestor_id"),
                    url=url,
                    extra_params=extra_attrs,
                    **optional,
                )
                reversion.set_comment(
                    reversion_comment or "Created by logic.data_import.import_sushi_credentials"
                )
                db_credentials[key] = cr
                logger.info('#%03d: Credentials created for platform "%s"', i + 2, platform.name_en)
            stats["added"] += 1
        # link report types
        linked_rts = {rt.code for rt in cr.counter_reports.all()}
        report_types = record.get("counter_reports", "").split(",")

        # Get platform provider for whitelisting checks
        provider = None
        if platform.knowledgebase:
            provider = get_provider_for_counter_version(platform.knowledgebase, version)
        for report_type in report_types:
            report_type = report_type.strip()
            if report_type and report_type not in linked_rts:
                if rt := CounterReportType.objects.filter(
                    code=report_type, counter_version=version
                ).first():
                    # Check whitelisting for reports that require it
                    if rt.requires_whitelisting:
                        if not provider:
                            logger.warning(
                                '#%03d: Report type "%s" requires whitelisting but platform has no '
                                "knowledgebase provider",
                                i + 2,
                                report_type,
                            )
                            stats["report_type_not_whitelisted"] += 1
                            continue
                        elif not is_report_type_whitelisted(provider, rt.code):
                            logger.warning(
                                '#%03d: Report type "%s" is not whitelisted for platform',
                                i + 2,
                                report_type,
                            )
                            stats["report_type_not_whitelisted"] += 1
                            continue

                    CounterReportsToCredentials.objects.create(credentials=cr, counter_report=rt)
                    stats["report_type_assigned"] += 1
                else:
                    logger.error('#%03d: Report type "%s" not found', i + 2, report_type)
                    stats["report_type_not_found"] += 1
    return stats


def parse_params(text, version: Optional[int] = None) -> dict:
    out = {}
    text = text.strip()
    for part in text.split(";"):
        if "=" in part:
            name, value = part.split("=")
            name = name.strip()
            value = value.strip()
            if name == "auth":
                value = tuple(value.split(","))
            out[name] = value
    if text and not out and version == 5:
        # there is some text, but we could not extract anything from it
        # if this is C5, we assume the value is the API key
        out["api_key"] = text
    return out


def get_int_value(mapping, key):
    value = mapping.get(key)
    if isinstance(value, int):
        return value
    if value is not None:
        return int(value.strip())
    return None
