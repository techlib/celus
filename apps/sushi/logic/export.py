import logging
from enum import Enum
from io import BytesIO
from typing import List, Optional

import pandas as pd
from core.models import DATA_SOURCE_TYPE_ORGANIZATION
from django.conf import settings
from django.db.models import QuerySet
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from organizations.models import Organization
from publications.models import Platform
from sushi.models import CounterReportType, SushiCredentials

logger = logging.getLogger(__name__)


class Col(Enum):
    TITLE = 'title'
    ORGANIZATION = 'organization'
    PUBLISHER_VENDOR_PLATFORM = 'publisher/vendor/platform'
    SUSHI_URL = 'SUSHI url'
    REQUESTOR_ID = 'requestor id'
    CUSTOMER_ID = 'customer id'
    API_KEY = 'api key'
    PLATFORM_FILTER = 'platform filter'
    HTTP_USERNAME = 'http username'
    HTTP_PASSWORD = 'http password'
    EXTRA_PARAMS = 'extra params'


class CredentialsDataFrame:
    def __init__(
        self,
        counter_version: int,
        cols: list,
        report_types: Optional[QuerySet[CounterReportType]] = None,
    ):
        self.counter_version = counter_version
        self.report_types = report_types
        self.cols = cols

    @classmethod
    def export(cls, counter_version: int):
        cols = [
            Col.TITLE,
            Col.ORGANIZATION,
            Col.PUBLISHER_VENDOR_PLATFORM,
            Col.SUSHI_URL,
            Col.REQUESTOR_ID,
            Col.CUSTOMER_ID,
        ]
        if counter_version == 5:
            cols += [Col.API_KEY, Col.PLATFORM_FILTER]
        if counter_version == 4:
            cols += [Col.HTTP_USERNAME, Col.HTTP_PASSWORD, Col.EXTRA_PARAMS]
        report_types = CounterReportType.objects.filter(
            counter_version=counter_version
        ).values_list('code', flat=True)
        report_type_cols = [rep_type for rep_type in report_types]
        cols += report_type_cols
        return cls(counter_version, cols, report_types)

    @classmethod
    def template_for_import(cls, selected_organization_id):
        cols = [
            Col.TITLE,
            Col.ORGANIZATION,
            Col.PUBLISHER_VENDOR_PLATFORM,
            Col.REQUESTOR_ID,
            Col.CUSTOMER_ID,
            Col.API_KEY,
            Col.PLATFORM_FILTER,
        ]
        if selected_organization_id != '-1':
            cols.remove(Col.ORGANIZATION)
        return cls(5, cols)

    def platforms_to_display(self, accessible_organizations) -> QuerySet[Platform]:
        platforms = Platform.objects.exclude(source__type=DATA_SOURCE_TYPE_ORGANIZATION)
        if accessible_organizations is not None:
            platforms = platforms | Platform.objects.filter(
                source__organization__in=accessible_organizations
            )
        return platforms.distinct().order_by('name_en')

    def create(
        self,
        sushi_credentials: QuerySet[SushiCredentials],
        accessible_organizations: Optional[QuerySet[Organization]] = None,
    ) -> pd.DataFrame:
        sushicred_dict = {col: [] for col in self.cols}
        credentials = sushi_credentials.filter(
            counter_version=self.counter_version
        ).prefetch_related('counter_reports__report_type')
        for cr in credentials:
            sushicred_dict[Col.TITLE].append(cr.title)
            if Col.ORGANIZATION in sushicred_dict:
                sushicred_dict[Col.ORGANIZATION].append(cr.organization.name_en)
            sushicred_dict[Col.PUBLISHER_VENDOR_PLATFORM].append(cr.platform.name_en)
            if Col.SUSHI_URL in sushicred_dict:
                sushicred_dict[Col.SUSHI_URL].append(cr.url)
            sushicred_dict[Col.REQUESTOR_ID].append(cr.requestor_id)
            sushicred_dict[Col.CUSTOMER_ID].append(cr.customer_id)

            report_types = {rep.report_type.short_name for rep in cr.counter_reports.all()}
            if self.report_types:
                for crt_code in self.report_types:
                    sushicred_dict[crt_code].append("active" if crt_code in report_types else "")
            if Col.API_KEY in sushicred_dict:
                sushicred_dict[Col.API_KEY].append(cr.api_key)
            if Col.PLATFORM_FILTER in sushicred_dict:
                sushicred_dict[Col.PLATFORM_FILTER].append(
                    cr.extra_params['platform'] if 'platform' in cr.extra_params else ''
                )
            if Col.HTTP_USERNAME in sushicred_dict:
                sushicred_dict[Col.HTTP_USERNAME].append(cr.http_username)
            if Col.HTTP_PASSWORD in sushicred_dict:
                sushicred_dict[Col.HTTP_PASSWORD].append(cr.http_password)
            if Col.EXTRA_PARAMS in sushicred_dict:
                sushicred_dict[Col.EXTRA_PARAMS].append(cr.extra_params)
        if accessible_organizations:
            platforms_in_sushicred_dict = set(sushicred_dict[Col.PUBLISHER_VENDOR_PLATFORM])
            missing_platforms = {
                p.name_en for p in self.platforms_to_display(accessible_organizations)
            } - platforms_in_sushicred_dict
            sushicred_dict[Col.PUBLISHER_VENDOR_PLATFORM] += list(missing_platforms)
            for item in missing_platforms:
                for col in sushicred_dict:
                    if col != Col.PUBLISHER_VENDOR_PLATFORM:
                        sushicred_dict[col].append("")
        sushicred_dict = {
            (k.value if isinstance(k, Enum) else k): v for k, v in sushicred_dict.items()
        }
        df = pd.DataFrame(sushicred_dict)
        return self.sort(df)

    def sort(self, df) -> pd.DataFrame:
        to_sort_by = [Col.PUBLISHER_VENDOR_PLATFORM.value, Col.TITLE.value]
        if Col.ORGANIZATION.value in df:
            to_sort_by.insert(1, Col.ORGANIZATION.value)
        return df.sort_values(by=to_sort_by)


class OrganizationsDataFrame:
    class Col(Enum):
        ORGANIZATION = 'organization'

    def __init__(self, admin_organizations: QuerySet[Organization]):
        self.admin_organizations = admin_organizations

    def create(self) -> pd.DataFrame:
        return pd.DataFrame(
            {
                Col.ORGANIZATION.value: [
                    organization.name_en for organization in self.admin_organizations
                ]
            }
        )


class Sheet:
    def __init__(self, df: pd.DataFrame, title: str):
        self.df = df
        self.title = title


class XlsxFile:
    def __init__(
        self, file: BytesIO, sheets: List[Sheet], selected_organization_id: Optional[str] = None
    ):
        self.file = file
        self.sheets = sheets
        self.selected_organization_id = selected_organization_id

    @staticmethod
    def adjust_col_widths(worksheet, df: pd.DataFrame):
        for idx, col in enumerate(df):
            series = df[col]
            max_len = max((series.astype(str).map(len).max(), len(str(series.name)))) + 0.5
            max_len = min(max_len, 60)
            worksheet.column_dimensions[get_column_letter(idx + 1)].width = max_len

    @classmethod
    def new(cls, sheets: List[Sheet]):
        return cls(BytesIO(), sheets)

    @classmethod
    def use_template(cls, sheets: List[Sheet], selected_organization_id: str):
        tmp_file = BytesIO()
        template_file = (
            settings.TEMPLATE_FOR_SUSHI_CRED_IMPORT_CONSORTIUM
            if selected_organization_id == '-1'
            else settings.TEMPLATE_FOR_SUSHI_CRED_IMPORT_SINGLE_ORG
        )
        openpyxl_workbook = load_workbook(template_file)
        openpyxl_workbook.save(tmp_file)
        return cls(tmp_file, sheets, selected_organization_id)

    def create(self, mode='w', if_sheet_exists: Optional[str] = None) -> BytesIO:
        with pd.ExcelWriter(
            self.file, engine='openpyxl', mode=mode, if_sheet_exists=if_sheet_exists
        ) as writer:
            for sheet in self.sheets:
                if sheet.title:
                    sheet.df.to_excel(writer, sheet_name=sheet.title, index=False)
                    XlsxFile.adjust_col_widths(writer.sheets[sheet.title], sheet.df)
        return self.file
